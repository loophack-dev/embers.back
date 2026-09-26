"""Builds .docx, .pptx and .md files from their specs. Pure functions returning bytes."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation

from embers.artifacts.specs import DocumentSpec, MarkdownSpec, PresentationSpec

MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "md": "text/markdown; charset=utf-8",
}

_TITLE_LAYOUT = 0
_CONTENT_LAYOUT = 1


def build_docx(spec: DocumentSpec) -> bytes:
    document = Document()
    document.add_heading(spec.title, level=0)
    if spec.summary:
        document.add_paragraph(spec.summary)
    for section in spec.sections:
        document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        for bullet in section.bullets:
            document.add_paragraph(bullet, style="List Bullet")
        if section.table is not None:
            columns = len(section.table.headers)
            table = document.add_table(rows=1, cols=columns)
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, section.table.headers, strict=True):
                cell.text = header
            for values in section.table.rows:
                cells = table.add_row().cells
                for index in range(columns):
                    cells[index].text = values[index] if index < len(values) else ""
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx(spec: PresentationSpec, template_path: Path | None = None) -> bytes:
    presentation = Presentation(str(template_path)) if template_path else Presentation()
    layouts = presentation.slide_layouts

    def layout(index: int) -> Any:
        # Custom templates may not have the standard layouts; fall back to the first one.
        return layouts[index] if index < len(layouts) else layouts[0]

    cover = presentation.slides.add_slide(layout(_TITLE_LAYOUT))
    if cover.shapes.title is not None:
        cover.shapes.title.text = spec.title
    if spec.subtitle and len(cover.placeholders) > 1:
        cover.placeholders[1].text = spec.subtitle

    for item in spec.slides:
        slide = presentation.slides.add_slide(layout(_CONTENT_LAYOUT))
        if slide.shapes.title is not None:
            slide.shapes.title.text = item.title
        body = next(
            (p for p in slide.placeholders if p.placeholder_format.idx != 0 and p.has_text_frame),
            None,
        )
        if body is not None and item.bullets:
            frame = body.text_frame
            frame.text = item.bullets[0]
            for bullet in item.bullets[1:]:
                frame.add_paragraph().text = bullet
        if item.notes:
            slide.notes_slide.notes_text_frame.text = item.notes

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def build_markdown(spec: MarkdownSpec) -> bytes:
    return f"# {spec.title}\n\n{spec.content.strip()}\n".encode()
